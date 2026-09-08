"""Report rendering (P5): one `Report` -> JSON, HTML/PDF (WeasyPrint), DOCX (python-docx).

The report invents nothing. It reads the declarations, violations and score the pipeline already
produced and rearranges them for a reader, so the three files can never disagree with each other
or with the scan page. `Report.report.schema_version` is what an API consumer pins against.
"""

from __future__ import annotations

import base64
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import cv2
import numpy as np
from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup
from numpy.typing import NDArray

from .models import (
    CheckStatus,
    Declaration,
    PipelineResult,
    Report,
    ReportBox,
    ReportCheck,
    ReportDeclaration,
    ReportFinding,
    ReportImage,
    ReportMeta,
    ReportScan,
    ReportScore,
    ScanContext,
    ScanRow,
    Severity,
    Violation,
    Word,
)
from .rules_engine import PENALTY, applicable_rules, load_rules

TEMPLATES = Path(__file__).resolve().parent / "templates"
ASSETS = TEMPLATES / "assets"

# The canonical field names of CLAUDE.md, written out. The report is read by an inspector and a
# shopkeeper, not by the person who chose the identifiers.
FIELD_LABELS: dict[str, str] = {
    "manufacturer": "Manufacturer / packer",
    "importer": "Importer",
    "generic_name": "Common or generic name",
    "net_quantity": "Net quantity",
    "mfg_date": "Month and year of manufacture",
    "mrp": "Retail sale price (MRP)",
    "consumer_care": "Consumer care details",
    "country_of_origin": "Country of origin",
    "unit_sale_price": "Unit sale price",
    "best_before": "Best before / use by",
}

STATUS_WORDS: dict[str, str] = {
    "compliant": "Compliant",
    "minor_issues": "Minor issues",
    "non_compliant": "Not compliant",
}

SCALE_WORDS: dict[str, str] = {
    "aruco": "ArUco marker in the frame",
    "card": "reference card in the frame",
    "inspector": "panel width entered by the inspector",
    "none": "no scale reference in the frame",
}


# ---------------------------------------------------------------- building the report


def _boxes(word_ids: list[int], words: dict[int, Word]) -> list[ReportBox]:
    return [
        ReportBox(image_id=w.image_id, x=w.x, y=w.y, w=w.w, h=w.h, text=w.text)
        for i in word_ids
        if (w := words.get(i)) is not None
    ]


def _finding(v: Violation, title: str, words: dict[int, Word]) -> ReportFinding:
    scored = v.evidence.status == CheckStatus.fail and v.severity != Severity.info
    return ReportFinding(
        rule_id=v.rule_id,
        rule_ref=v.rule_ref,
        title=title,
        severity=v.severity,
        status=v.evidence.status,
        message=v.message,
        reason=v.evidence.reason,
        values=v.evidence.values,
        boxes=_boxes(v.evidence.word_ids, words),
        confidence=v.confidence,
        penalty=PENALTY[v.severity] if scored else 0,
    )


def _declaration(d: Declaration, words: dict[int, Word]) -> ReportDeclaration:
    return ReportDeclaration(
        field=d.field,
        label=FIELD_LABELS.get(d.field, d.field.replace("_", " ").capitalize()),
        value=d.value,
        confidence=d.confidence,
        height_mm=d.height_mm,
        width_height_ratio=d.width_height_ratio,
        boxes=_boxes(d.word_ids, words),
    )


def build_report(
    scan: ScanRow,
    result: PipelineResult,
    ctx: ScanContext,
    images: list[ReportImage] | None = None,
    now: datetime | None = None,
) -> Report:
    """Rearrange a finished scan into the shape the three files are rendered from."""
    words = {w.id: w for w in result.words}
    rules = load_rules()
    titles = {r.rule_id: r.title for r in rules}

    violations: list[ReportFinding] = []
    notes: list[ReportFinding] = []
    unverifiable: list[ReportFinding] = []
    for v in result.violations:
        finding = _finding(v, titles.get(v.rule_id, v.rule_id), words)
        if finding.status == CheckStatus.unverifiable:
            unverifiable.append(finding)
        elif finding.severity == Severity.info:
            notes.append(finding)
        else:
            violations.append(finding)

    # A rule that raised nothing passed — but only rules that can fail. An `info` rule is a flag
    # (F3 medical device, E2, D9 best-before) or an exemption note (X1..X4): it says something
    # when it fires and nothing when it does not, so "F3 passed — medical device pack" would be
    # a sentence with no meaning in it.
    raised = {v.rule_id for v in result.violations}
    passed = [
        ReportCheck(rule_id=r.rule_id, rule_ref=r.rule_ref, title=r.title)
        for r in applicable_rules(result.declarations, ctx, rules)
        if r.rule_id not in raised and r.severity != Severity.info
    ]

    counts = {s: sum(1 for v in violations if v.severity == s) for s in Severity}
    status = (
        "non_compliant"
        if counts[Severity.critical] or counts[Severity.major]
        else "minor_issues"
        if counts[Severity.minor]
        else "compliant"
    )

    return Report(
        report=ReportMeta(generated_at=now or datetime.now(UTC)),
        scan=ReportScan(
            id=scan.id,
            source=scan.source,
            created_at=scan.created_at,
            notes=scan.notes,
            has_reference_card=scan.has_reference_card,
            mm_per_px=result.mm_per_px,
            scale_source=result.scale_source,
            pdp_area_cm2=ctx.pdp_area_cm2,
        ),
        score=ReportScore(
            value=result.compliance_score,
            status=cast(Any, status),
            critical=counts[Severity.critical],
            major=counts[Severity.major],
            minor=counts[Severity.minor],
            notes=len(notes),
            unverifiable=len(unverifiable),
            penalty=sum(v.penalty for v in violations),
        ),
        declarations=[_declaration(d, words) for d in result.declarations],
        violations=violations,
        notes=notes,
        unverifiable=unverifiable,
        passed=passed,
        images=images or [],
    )


# ---------------------------------------------------------------- evidence photographs


def annotate(
    image: NDArray[np.uint8],
    declarations: list[ReportBox],
    evidence: list[ReportBox],
    max_width: int = 900,
) -> bytes:
    """The photograph with a box round every declaration, and the cited evidence in red.

    Boxes are burned into the pixels because the same JPEG goes into the PDF and the DOCX, and
    a Word document cannot position an overlay. Downscaled first: a report an inspector opens
    on a phone should not carry three 1600 px photographs.
    """
    scale = min(1.0, max_width / max(1, image.shape[1]))
    out = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
    cited = {(b.x, b.y, b.w, b.h) for b in evidence}
    for box in [*declarations, *evidence]:
        colour = (0, 0, 210) if (box.x, box.y, box.w, box.h) in cited else (190, 120, 0)
        x, y = int(box.x * scale), int(box.y * scale)
        cv2.rectangle(out, (x, y), (x + int(box.w * scale), y + int(box.h * scale)), colour, 2)
    ok, buf = cv2.imencode(".jpg", out, [int(cv2.IMWRITE_JPEG_QUALITY), 78])
    if not ok:
        raise ValueError("could not encode the evidence photograph")
    return buf.tobytes()


def annotate_panels(paths: dict[str, Path], report: Report) -> dict[str, bytes]:
    """One annotated JPEG per photograph, keyed by scan_images.id."""
    evidence = [b for v in report.violations for b in v.boxes]
    declarations = [b for d in report.declarations for b in d.boxes]
    panels: dict[str, bytes] = {}
    for image_id, path in paths.items():
        img = cv2.imread(str(path))
        if img is None:
            continue
        panels[image_id] = annotate(
            cast("NDArray[np.uint8]", img),
            [b for b in declarations if b.image_id == image_id],
            [b for b in evidence if b.image_id == image_id],
        )
    return panels


# ---------------------------------------------------------------- rendering


def render_json(report: Report) -> str:
    return report.model_dump_json(indent=2)


def _icon(name: str) -> Markup:
    """A Lucide icon inlined. Vendored under templates/assets/icons (ISC) so the report never
    reaches out to a website while it renders. `Markup` because autoescaping is on for
    everything the pack itself printed, and an escaped `<svg>` is a paragraph of angle
    brackets — these five files are ours, not input."""
    return Markup((ASSETS / "icons" / f"{name}.svg").read_text(encoding="utf-8"))  # noqa: S704


def _environment() -> Environment:
    env = Environment(
        loader=FileSystemLoader(TEMPLATES),
        autoescape=select_autoescape(["html"]),
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.globals["icon"] = _icon
    env.globals["STATUS_WORDS"] = STATUS_WORDS
    env.globals["SCALE_WORDS"] = SCALE_WORDS
    return env


def render_html(report: Report, panels: dict[str, bytes] | None = None) -> str:
    """The report as one self-contained HTML page. Photographs are inlined as data URIs so the
    same string renders identically in a browser and inside WeasyPrint."""
    photos = [
        {
            "kind": image.kind,
            "src": "data:image/jpeg;base64," + base64.b64encode(panels[image.id]).decode("ascii"),
        }
        for image in report.images
        if panels and image.id in panels
    ]
    return (
        _environment()
        .get_template("report.html")
        .render(r=report, photos=photos, summary=_summary_rows(report), footnotes=FOOTNOTES)
    )


def render_pdf(html: str) -> bytes:
    """WeasyPrint. Imported here and not at the top of the file: it binds to Pango and Cairo at
    import time, which a machine without the GTK libraries does not have, and a missing PDF must
    not stop the JSON and the DOCX from being written."""
    from weasyprint import HTML  # noqa: PLC0415

    return cast(bytes, HTML(string=html, base_url=str(TEMPLATES)).write_pdf())


# ---------------------------------------------------------------- DOCX

_SEVERITY_WORD: dict[Severity, str] = {
    Severity.critical: "CRITICAL",
    Severity.major: "MAJOR",
    Severity.minor: "MINOR",
    Severity.info: "INFO",
}


def render_docx(report: Report, panels: dict[str, bytes] | None = None) -> bytes:
    """The same report as a Word document: title, summary, score, violations, declarations,
    evidence, notes. Same hierarchy as the PDF, no styling tricks that Word would drop."""
    from docx import Document  # noqa: PLC0415
    from docx.shared import Inches, Pt  # noqa: PLC0415

    doc = Document()
    doc.core_properties.title = f"Legal Metrology compliance report — scan {report.scan.id}"
    doc.add_heading("Legal Metrology compliance report", level=0)
    doc.add_paragraph(report.report.law)

    doc.add_heading("Summary", level=1)
    summary = doc.add_table(rows=0, cols=2)
    summary.style = "Table Grid"
    for name, value in _summary_rows(report):
        cells = summary.add_row().cells
        cells[0].text = name
        cells[1].text = value

    verdict = doc.add_paragraph()
    run = verdict.add_run(
        f"{STATUS_WORDS[report.score.status]} — {report.score.value} out of {report.score.out_of}"
    )
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph(
        f"{report.score.critical} critical, {report.score.major} major, "
        f"{report.score.minor} minor. {report.score.unverifiable} check(s) could not be "
        f"verified and cost no points."
    )

    doc.add_heading("Violations", level=1)
    if not report.violations:
        doc.add_paragraph("No violation was found.")
    for v in report.violations:
        doc.add_heading(f"{v.rule_id} · {v.rule_ref} · {_SEVERITY_WORD[v.severity]}", level=2)
        doc.add_paragraph(v.title)
        doc.add_paragraph(v.message)
        doc.add_paragraph(f"Score: −{v.penalty} points.")
        for key, value in v.values.items():
            doc.add_paragraph(f"{key}: {value}", style="List Bullet")
        if v.boxes:
            doc.add_paragraph(
                "Evidence on the photograph: "
                + "; ".join(f'"{b.text}"' for b in v.boxes if b.text),
                style="List Bullet",
            )
        elif v.status == CheckStatus.fail:
            doc.add_paragraph(
                "No evidence box: the declaration is not on the photographs.", style="List Bullet"
            )

    doc.add_heading("Not verifiable (for information)", level=1)
    if not report.unverifiable:
        doc.add_paragraph("Every applicable check could be verified.")
    for v in report.unverifiable:
        doc.add_paragraph(
            f"{v.rule_id} · {v.rule_ref} — {v.title}. {v.reason or v.message}",
            style="List Bullet",
        )

    doc.add_heading("Declarations read from the package", level=1)
    if not report.declarations:
        doc.add_paragraph("No declaration was read from the photographs.")
    else:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, head in zip(
            table.rows[0].cells, ("Declaration", "Value", "Print height"), strict=True
        ):
            cell.text = head
        for d in report.declarations:
            cells = table.add_row().cells
            cells[0].text = d.label
            cells[1].text = d.value
            cells[2].text = "not measured" if d.height_mm is None else f"{d.height_mm:.1f} mm"

    doc.add_heading("Checks that passed", level=1)
    if not report.passed:
        doc.add_paragraph("No check passed.")
    for check in report.passed:
        doc.add_paragraph(
            f"{check.rule_id} · {check.rule_ref} — {check.title}", style="List Bullet"
        )

    if report.notes:
        doc.add_heading("Notes", level=1)
        for note in report.notes:
            doc.add_paragraph(
                f"{note.rule_id} · {note.rule_ref} — {note.message}", style="List Bullet"
            )

    if panels:
        doc.add_heading("Evidence photographs", level=1)
        doc.add_paragraph(
            "Blue: a declaration the checker read. Red: evidence cited by a violation above."
        )
        for image in report.images:
            if image.id in panels:
                doc.add_picture(BytesIO(panels[image.id]), width=Inches(6))
                doc.add_paragraph(image.kind)

    doc.add_heading("About this report", level=1)
    for line in FOOTNOTES:
        doc.add_paragraph(line)
    doc.add_paragraph(
        f"Generated {report.report.generated_at:%Y-%m-%d %H:%M UTC} by the Legal Metrology "
        "compliance checker. Schema " + report.report.schema_version + "."
    )

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


FOOTNOTES: list[str] = [
    "Penalties under Section 36 of the Legal Metrology Act, 2009: up to ₹25,000 for a first "
    "offence, ₹50,000 for a second, ₹1,00,000 and possible imprisonment thereafter. Net "
    "quantity errors: ₹10,000–50,000.",
    "A check marked “not verifiable” was not measured, because the photographs did not carry "
    "what the measurement needs. It costs no points. Nothing in this report is estimated.",
    "This report describes the photographs that were uploaded. A declaration printed on a panel "
    "nobody photographed is reported as absent.",
]


def _summary_rows(report: Report) -> list[tuple[str, str]]:
    """The header facts, in one place so the PDF and the DOCX cannot drift apart."""
    scan = report.scan
    scale = SCALE_WORDS[str(scan.scale_source)]
    return [
        ("Scan ID", scan.id),
        (
            "Scanned",
            scan.created_at.strftime("%d %b %Y, %H:%M") if scan.created_at else "—",
        ),
        ("Source", "E-commerce listing" if scan.source == "ecommerce" else "Package photograph"),
        (
            "Print scale",
            f"1 pixel = {scan.mm_per_px:.4f} mm ({scale})"
            if scan.mm_per_px
            else f"not measured — {scale}",
        ),
        (
            "Principal display panel",
            f"{scan.pdp_area_cm2:.1f} cm²" if scan.pdp_area_cm2 else "not given",
        ),
        ("Inspector's notes", scan.notes or "—"),
    ]
